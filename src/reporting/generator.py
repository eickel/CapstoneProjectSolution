"""Multi-format report generator & exporter (Excel, CSV, JSON)."""

import json
from datetime import date
from pathlib import Path

import pandas as pd
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from src.config import settings
from src.exceptions import ReportGenerationError
from src.processing.analytics import DataProcessor


class ReportGenerator:
    """Builds the Excel workbook, CSV summary, and JSON export from a DataProcessor."""

    def __init__(self, processor: DataProcessor, output_dir: str | None = None):
        self.processor = processor
        self.output_dir = Path(output_dir or settings.reports_output_dir)
        self.output_dir.mkdir(parents=True, exist_ok=True)

    def _write_sheet(self, writer: pd.ExcelWriter, df: pd.DataFrame, sheet_name: str) -> None:
        df.to_excel(writer, sheet_name=sheet_name, index=False)
        worksheet = writer.sheets[sheet_name]
        for col_idx, column in enumerate(df.columns, start=1):
            worksheet.cell(row=1, column=col_idx).font = Font(bold=True)
            content_len = df[column].astype(str).str.len().max() if len(df) else 0
            max_len = max(len(str(column)), content_len)
            worksheet.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 2, 50)

    def generate_excel_report(self) -> Path:
        try:
            timestamp = date.today().strftime("%Y%m%d")
            path = self.output_dir / f"library_report_{timestamp}.xlsx"
            with pd.ExcelWriter(path, engine="openpyxl") as writer:
                self._write_sheet(writer, self.processor.subject_aggregation(), "Subjects")
                self._write_sheet(writer, self.processor.top_publication_years(), "Top Years")
                self._write_sheet(writer, self.processor.top_languages(), "Top Languages")
                self._write_sheet(writer, self.processor.author_productivity(), "Top Authors")
                summary_df = pd.DataFrame([self.processor.summary_metrics()])
                self._write_sheet(writer, summary_df, "Summary")
            return path
        except Exception as exc:
            raise ReportGenerationError(f"Failed to generate Excel report: {exc}") from exc

    def generate_csv_summary(self) -> Path:
        try:
            path = self.output_dir / "top_years_languages_summary.csv"
            years = self.processor.top_publication_years().assign(metric="publication_year")
            years = years.rename(columns={"first_publish_year": "value"})
            langs = self.processor.top_languages().assign(metric="language")
            langs = langs.rename(columns={"language": "value"})
            columns = ["metric", "value", "work_count"]
            combined = pd.concat([years, langs], ignore_index=True)[columns]
            combined.to_csv(path, index=False)
            return path
        except Exception as exc:
            raise ReportGenerationError(f"Failed to generate CSV summary: {exc}") from exc

    def generate_json_dump(self) -> Path:
        try:
            path = self.output_dir / "library_report.json"
            payload = {
                "summary": self.processor.summary_metrics(),
                "subjects": self.processor.subject_aggregation().to_dict(orient="records"),
                "top_years": self.processor.top_publication_years().to_dict(orient="records"),
                "top_languages": self.processor.top_languages().to_dict(orient="records"),
                "top_authors": self.processor.author_productivity().to_dict(orient="records"),
            }
            serialized = json.dumps(payload, separators=(",", ":"), default=str)
            path.write_text(serialized, encoding="utf-8")
            return path
        except Exception as exc:
            raise ReportGenerationError(f"Failed to generate JSON dump: {exc}") from exc

    def generate_all(self) -> dict:
        return {
            "excel": self.generate_excel_report(),
            "csv": self.generate_csv_summary(),
            "json": self.generate_json_dump(),
        }
